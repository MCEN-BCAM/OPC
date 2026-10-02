
from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any
import datetime as dt

import xlrd
from openpyxl import load_workbook

from .utils import unique_headers, json_safe


@dataclass
class SheetData:
    name: str
    headers: list[str]
    original_headers: list[Any]
    rows: list[dict[str, Any]]
    nrows: int
    ncols: int

    def to_metadata(self, sample_rows: int = 3) -> dict:
        return {
            "name": self.name,
            "nrows": self.nrows,
            "ncols": self.ncols,
            "original_headers": [json_safe(x) for x in self.original_headers],
            "headers": self.headers,
            "sample_rows": self.rows[:sample_rows],
        }


@dataclass
class WorkbookData:
    path: str
    sheets: dict[str, SheetData]

    def metadata(self) -> dict:
        return {
            "path": self.path,
            "sheet_names": list(self.sheets),
            "sheets": {
                name: sheet.to_metadata()
                for name, sheet in self.sheets.items()
            },
        }


def _convert_xlrd_cell(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> Any:
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
        except Exception:
            return cell.value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        value = float(cell.value)
        return int(value) if value.is_integer() else value
    return json_safe(cell.value)


def read_xls(path: str | Path) -> WorkbookData:
    path = Path(path)
    book = xlrd.open_workbook(path)
    sheets: dict[str, SheetData] = {}
    for ws in book.sheets():
        if ws.nrows == 0:
            headers: list[str] = []
            originals: list[Any] = []
            rows: list[dict[str, Any]] = []
        else:
            originals = [_convert_xlrd_cell(book, ws.cell(0, c)) for c in range(ws.ncols)]
            headers = unique_headers(originals)
            rows = []
            for r in range(1, ws.nrows):
                record = {
                    headers[c]: _convert_xlrd_cell(book, ws.cell(r, c))
                    for c in range(ws.ncols)
                }
                if any(v not in (None, "") for v in record.values()):
                    rows.append(record)
        sheets[ws.name] = SheetData(
            name=ws.name,
            headers=headers,
            original_headers=originals,
            rows=rows,
            nrows=ws.nrows,
            ncols=ws.ncols,
        )
    return WorkbookData(str(path), sheets)


def read_xlsx(path: str | Path) -> WorkbookData:
    path = Path(path)
    book = load_workbook(path, read_only=True, data_only=True)
    sheets: dict[str, SheetData] = {}
    try:
        for ws in book.worksheets:
            values = list(ws.iter_rows(values_only=True))
            if not values:
                originals: list[Any] = []
                headers: list[str] = []
                rows: list[dict[str, Any]] = []
            else:
                originals = [json_safe(v) for v in values[0]]
                headers = unique_headers(originals)
                rows = []
                for row in values[1:]:
                    padded = list(row) + [None] * (len(headers) - len(row))
                    record = {
                        headers[c]: json_safe(padded[c])
                        for c in range(len(headers))
                    }
                    if any(v not in (None, "") for v in record.values()):
                        rows.append(record)
            sheets[ws.title] = SheetData(
                name=ws.title,
                headers=headers,
                original_headers=originals,
                rows=rows,
                nrows=ws.max_row,
                ncols=ws.max_column,
            )
    finally:
        book.close()
    return WorkbookData(str(path), sheets)


def read_workbook(path: str | Path) -> WorkbookData:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return read_xls(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    raise ValueError(f"Unsupported spreadsheet format: {path}")
